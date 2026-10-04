"""Generate VirtualDev Studio professional presentation (PPTX)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

DARK = RGBColor(0x0F, 0x0F, 0x23)
CARD = RGBColor(0x1A, 0x1A, 0x2E)
ACCENT = RGBColor(0x00, 0xD4, 0xAA)
PURPLE = RGBColor(0x7C, 0x3A, 0xE0)
WHITE = RGBColor(0xE0, 0xE0, 0xE0)
GREY = RGBColor(0x9A, 0x9A, 0xB0)
RED = RGBColor(0xFF, 0x6B, 0x6B)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_bg(slide, color=DARK):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def add_rect(slide, left, top, width, height, color=CARD, round_=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color
    return shape


def add_text(slide, left, top, width, height, text, size=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = "Calibri"
    return box


def add_bullets(slide, left, top, width, height, items, size=16, color=WHITE, spacing=8):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = f"• {item}"
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = "Calibri"
    return box


def title_slide():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_rect(slide, 0, 0, 13.333, 0.15, ACCENT, round_=False)
    add_text(slide, 1.2, 2.2, 11, 1.0, "VirtualDev Studio", 44, WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 1.2, 3.2, 11, 0.6, "AI-Powered Virtual Software Development Studio", 22, ACCENT, align=PP_ALIGN.CENTER)
    add_text(slide, 1.2, 4.0, 11, 0.5, "An end-to-end platform that transforms ideas into production-ready software projects", 16, GREY, align=PP_ALIGN.CENTER)
    add_text(slide, 1.2, 6.5, 11, 0.4, "Product Team + Engineering Team + AI Team + QA + Security + Documentation", 13, GREY, align=PP_ALIGN.CENTER)


def section_slide(title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_text(slide, 1.0, 2.8, 11, 0.8, title, 36, WHITE, bold=True, align=PP_ALIGN.CENTER)
    if subtitle:
        add_text(slide, 1.0, 3.7, 11, 0.5, subtitle, 18, ACCENT, align=PP_ALIGN.CENTER)


def content_slide(title, bullets, sub=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_text(slide, 0.7, 0.3, 12, 0.6, title, 26, ACCENT, bold=True)
    add_rect(slide, 0.7, 0.85, 3.5, 0.06, ACCENT)
    add_bullets(slide, 0.9, 1.2, 11.5, 5.5, bullets, size=16, spacing=10)
    if sub:
        add_text(slide, 0.9, 6.8, 11, 0.4, sub, 12, GREY)


# 1. Title
title_slide()

# 2. Problem
content_slide("The Problem", [
    "Software development requires coordination across many specialized roles",
    "Small teams and solo developers lack access to all required expertise",
    "Existing AI tools generate code snippets but not complete project deliverables",
    "Fragmented tooling causes context loss and quality gaps",
    "Security review, testing, and documentation are often skipped or rushed",
])

# 3. Existing approach
content_slide("Existing Approach", [
    "AI Coding Assistants: GitHub Copilot, Cursor, Codeium — code completion focus",
    "AI Development Agents: Devin, OpenHands — autonomous coding but limited orchestration",
    "Project Management Tools: Jira, Linear — track work but don't generate it",
    "AI Design Platforms: v0, Lovable — UI generation without full lifecycle",
    "Gap: No tool combines multi-role orchestration + human approval + full deliverables",
])

# 4. Proposed solution
content_slide("Proposed Solution: VirtualDev Studio", [
    "AI-powered virtual development studio with 10 specialized agents",
    "Coordinated by a Master AI Orchestrator",
    "Controlled workflow with human approval checkpoints",
    "Delivers complete project packages: requirements, architecture, code, tests",
    "Web dashboard + mobile companion app",
    "Modular architecture: React + FastAPI + PostgreSQL + LLM abstraction",
])

# 5. Target users
content_slide("Target Users", [
    "Solo developers and indie hackers — full virtual team on demand",
    "Non-technical startup founders — idea-to-specification translation",
    "Small development teams — accelerate secondary tasks",
    "Project managers — single dashboard for status and approvals",
    "Freelance developers — automated docs, research, and presentations",
])

# 6. Key features
content_slide("Key Features", [
    "Idea-to-project workflow with 14 stages",
    "Master AI Orchestrator with task planning and agent delegation",
    "Human approval gates for architecture, development, security",
    "Real-time status via WebSocket",
    "AI Team monitoring page",
    "Testing & security review stages",
    "Document generation and diagram creation",
    "Mobile companion with approval workflows",
])

# 7. System workflow
content_slide("System Workflow", [
    "User Idea → Requirements Analysis → Research → Product Analysis",
    "System Architecture → UI/UX Design → Database Design",
    "Implementation Plan → Development → Testing",
    "Security Review → Documentation → Diagrams → Presentation",
    "Human approval checkpoints at architecture, development, and security",
])

# 8. Architecture
content_slide("Architecture", [
    "Frontend: React + TypeScript — 15 pages, responsive design",
    "Backend: Python + FastAPI — REST + WebSocket APIs",
    "Database: PostgreSQL — 14 tables, full relational model",
    "AI Layer: Provider abstraction with mock LLM for MVP",
    "Mobile: Flutter companion app",
    "Communication: JWT auth, rate limiting, structured error envelopes",
])

# 9. AI Agent architecture
content_slide("AI Agent Architecture", [
    "Master Orchestrator coordinates the workflow",
    "10 specialized agents: Requirements, Researcher, Strategist, Architect",
    "Database Engineer, Designer, Developer, Tester, Security, Docs",
    "Task state persisted across the full lifecycle",
    "Failure detection with retry and escalation",
    "Tool layer: MCP-based with permission boundaries",
])

# 10. Database
content_slide("Database Design", [
    "PostgreSQL with 14 tables: users, projects, stages, requirements",
    "research_items, architecture_components, agents, tasks, task_runs",
    "documents, diagrams, presentations, approvals, audit_logs",
    "17 targeted indexes for query patterns",
    "Full referential integrity with cascade deletes",
    "Complete ERD in editable Mermaid format",
])

# 11. UI screenshots
content_slide("UI Overview", [
    "Landing page with value proposition and workflow",
    "Dashboard: active projects, progress, pending approvals",
    "Project workspace with stage timeline",
    "AI Team page with agent roles and status",
    "Requirements, Research, Architecture, Testing screens",
    "Dark professional theme with responsive layout",
])

# 12. Mobile application
content_slide("Mobile Application (Flutter)", [
    "View projects and their progress on the go",
    "Create projects with structured input",
    "Approve/reject workflows from mobile",
    "Agent status and task updates",
    "WebSocket real-time notifications",
    "Dark theme matching web app (teal/purple)",
])

# 13. Security
content_slide("Security", [
    "JWT authentication with refresh tokens",
    "Rate limiting on auth endpoints",
    "Input validation with Pydantic",
    "RBAC model with project ownership checks",
    "Secrets via environment variables only",
    "Comprehensive 12-area security review performed",
])

# 14. Testing
content_slide("Testing", [
    "25 API tests: auth, projects, entities, orchestration",
    "All 25 tests passing",
    "Frontend TypeScript compilation verified",
    "Flutter analyze passing",
    "Identified and fixed orchestration bug during testing",
])

# 15. Results
content_slide("Results", [
    "Working FastAPI backend with 25+ endpoints",
    "Complete React dashboard (15 pages)",
    "Flutter mobile app",
    "AI orchestrator with 10 agent roles",
    "Human approval workflow",
    "12 professional diagrams (Mermaid)",
    "Full documentation package",
])

# 16. Business model
content_slide("Business Model", [
    "SaaS subscription tiers — PROPOSAL",
    "Free tier: 1 project, basic features",
    "Pro tier $29-49/month: unlimited projects",
    "Team tier $99-199/month: collaboration",
    "Enterprise: custom deployment",
    "Pay-per-outcome or usage-based alternatives possible",
])

# 17. Competitive advantage
content_slide("Competitive Advantage", [
    "Multi-role orchestration — not just code generation",
    "Human approval gates — user stays in control",
    "Complete deliverable package — docs, diagrams, presentation",
    "Two interfaces: web dashboard + mobile companion",
    "Open LLM abstraction — not locked to one provider",
    "Production-oriented mindset: security + testing built-in",
])

# 18. Future roadmap
content_slide("Future Roadmap", [
    "Real LLM provider integration (OpenAI, Anthropic, local)",
    "Actual code generation and execution",
    "MCP tool layer with live filesystem, git, terminal",
    "Team collaboration and shared workspaces",
    "GitHub/GitLab integration",
    "Automated diagram rendering",
    "Multi-user RBAC refinement",
])

# 19. Conclusion
content_slide("Conclusion", [
    "VirtualDev Studio demonstrates a coordinated AI development team",
    "Validated the full software lifecycle approach",
    "Real implementation: backend, frontend, mobile, orchestrator, tests",
    "Human-in-the-loop design keeps users in control",
    "Ready for Phase 2: real LLM integration and tool layer",
])

# 20. Q&A
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_text(slide, 1.2, 3.0, 11, 1.0, "Thank You", 40, WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, 1.2, 4.0, 11, 0.6, "Questions & Discussion", 22, ACCENT, align=PP_ALIGN.CENTER)
add_text(slide, 1.2, 5.0, 11, 0.4, "VirtualDev Studio — Idea to Production, Orchestrated", 14, GREY, align=PP_ALIGN.CENTER)

from pathlib import Path

prs.save(Path(__file__).resolve().parent.parent / "docs/presentation/VirtualDev-Studio-Presentation.pptx")
print(f"Presentation saved: {len(prs.slides)} slides")